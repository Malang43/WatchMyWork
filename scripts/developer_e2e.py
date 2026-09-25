"""Full developer demo: real extension, live inference, export and Bob handoff."""
import io
import json
import re
import tempfile
import time
import zipfile
from pathlib import Path
import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
client = httpx.Client(base_url='http://127.0.0.1:8000', headers={'X-WatchMyWork': 'local-demo'}, timeout=400)
assert client.get('/teach/current').json() is None, 'Finish current recording first'
assert not any(r['state'] in ('running', 'paused') for r in client.get('/runs').json()), 'Finish current run first'
with sync_playwright() as pw:
    extension = str(ROOT / 'extension')
    context = pw.chromium.launch_persistent_context(tempfile.mkdtemp(prefix='wmw-developer-'), channel='chromium', headless=True, args=[f'--disable-extensions-except={extension}', f'--load-extension={extension}'], viewport={'width': 1440, 'height': 1000})
    try:
        page = context.new_page()
        errors = []
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.goto('http://127.0.0.1:5174/')
        page.get_by_role('button', name='New Workflow', exact=True).last.click()
        payload = (ROOT / 'sample-data/developer-tests.csv').read_bytes() + b'dev@example.com,test123,Deliberate mismatch,,\n'
        with page.expect_response(lambda r: r.url.endswith('/datasets') and r.request.method == 'POST') as response:
            page.get_by_label('Upload spreadsheet').set_input_files({'name': 'developer-e2e.csv', 'mimeType': 'text/csv', 'buffer': payload})
        dataset = response.value.json()
        page.get_by_label('Input 1', exact=True).select_option('Email')
        page.get_by_label('Input 2 (optional)', exact=True).select_option('Password')
        page.get_by_label('Destination / output column', exact=True).select_option('Actual Result')
        page.get_by_label('Expected output', exact=True).select_option('Expected Result')
        page.get_by_label('Test status', exact=True).select_option('Status')
        page.get_by_role('button', name='Continue to Teach').click()
        page.get_by_role('button', name='Start Recording').click()
        portal = context.new_page()
        portal.goto('http://127.0.0.1:5173/developer')
        overlay = portal.locator('#watchmywork-recorder')
        expect(overlay).to_have_attribute('data-recording', re.compile('[a-f0-9]{32}'), timeout=15000)
        portal.locator('#email-input').fill('dev@example.com')
        portal.locator('#password-input').fill('test123')
        portal.locator('#login-button').click()
        expect(overlay).to_have_text('✓ Test demonstration captured', timeout=15000)
        demo = client.get('/demos', params={'dataset_id': dataset['id']}).json()[0]
        assert demo['state'] == 'complete'
        assert len([e for e in demo['events'] if e['action'] == 'extract']) == 1
        print('Two-input recording and auto-save passed', flush=True)
        with page.expect_response(lambda r: r.url.endswith('/workflows/infer'), timeout=400000) as inference:
            page.get_by_role('button', name='Analyze Workflow').click()
        assert inference.value.status == 200, inference.value.text()
        workflow = inference.value.json()
        print('Inference source:', workflow['source'], flush=True)
        page.get_by_role('button', name='Confirm Workflow').click()
        page.get_by_role('button', name='Generate Playwright Test', exact=True).click()
        expect(page.locator('pre')).to_contain_text('for index, row in enumerate(cases, 1)')
        download = client.get(f"/workflows/{workflow['id']}/test-download")
        directory = ROOT / 'test-results' / 'generated-developer-test'
        directory.mkdir(exist_ok=True)
        with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
            archive.extractall(directory)
        calls = client.get('/metrics').json()['total_ai_calls']
        with page.expect_response(lambda r: r.url.endswith('/runs') and r.request.method == 'POST') as response:
            page.get_by_role('button', name='Run Test Suite').click()
        run = response.value.json()
        for _ in range(240):
            run = client.get(f"/runs/{run['id']}").json()
            if run['state'] == 'completed': break
            assert run['state'] == 'running', run['state']
            time.sleep(.25)
        assert run['state'] == 'completed'
        assert (run['PASS'], run['FAIL'], run['ERROR']) == (5, 1, 0), run
        assert client.get('/metrics').json()['total_ai_calls'] == calls
        report = client.get(f"/runs/{run['id']}/test-report").json()
        assert report[4]['actual'] == 'Email and password are required'
        assert report[5]['status'] == 'FAIL'
        assert client.get(f"/runs/{run['id']}/exceptions").json() == []
        book = load_workbook(io.BytesIO(client.get(f"/runs/{run['id']}/download").content))
        assert book.active.cell(7, 4).value == 'Login successful'
        assert book.active.cell(7, 5).value == 'FAIL'
        expect(page.get_by_role('heading', name='Test Report', exact=True)).to_be_visible()
        page.get_by_role('button', name='Prepare Bob Debug Package').click()
        expect(page.get_by_text('Package ready:', exact=False)).to_be_visible(timeout=15000)
        package = ROOT / 'bob_debug_package' / run['id']
        for name in ['generated_test.spec.py', 'test_cases.json', 'failed_tests.json', 'workflow.json', 'selectors.json', 'runtime_logs.json', 'failure_summary.md']:
            assert (package / name).is_file(), name
        assert len(json.loads((package / 'failed_tests.json').read_text())) == 1
        assert (ROOT / 'bob_sessions' / 'README.md').is_file()
        assert not errors, errors
        page.screenshot(path=str(ROOT / 'test-results' / 'developer-report.png'), full_page=True)
        print('Developer E2E passed: 6 executed, 5 PASS, 1 deliberate FAIL, 0 ERROR; report, export, generated test and Bob package verified.', flush=True)
    finally:
        context.close()

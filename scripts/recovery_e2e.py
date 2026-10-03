"""Real browser approval + memory demo. Run developer_e2e.py first; uses its synthetic workflow.

All three local servers must be running. No overlapping E2E sessions.
"""
import io
import json
import time
import zipfile
from pathlib import Path
import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'test-results'
ARTIFACTS.mkdir(exist_ok=True)
client = httpx.Client(base_url='http://127.0.0.1:8000', headers={'X-WatchMyWork': 'local-demo'}, timeout=30)
active = [r for r in client.get('/runs').json() if r['state'] in ('running', 'paused')]
assert not active or (len(active) == 1 and active[0].get('recovery_demo') and active[0]['state'] == 'paused'), 'Finish existing run first'
workflow = next(w for w in client.get('/workflows').json() if w['confirmed'] and w['plan']['mode'] == 'developer')


def wait_run(run_id):
    for _ in range(240):
        run = client.get(f'/runs/{run_id}').json()
        if run['state'] == 'completed': return run
        assert run['state'] == 'running', run['state']
        time.sleep(.5)
    raise AssertionError('Recovery run did not complete')


with sync_playwright() as pw:
    browser = pw.chromium.launch()
    page = browser.new_page(viewport={'width': 1440, 'height': 1100})
    try:
        portal = browser.new_page()
        # Query mode is deterministic, including after a prior changed-mode visit.
        for suffix, selector, name in [('', '#login-button', 'Login'), ('?recovery_demo=1', '#sign-in-button', 'Sign In'), ('', '#login-button', 'Login')]:
            portal.goto('http://127.0.0.1:5173/developer' + suffix)
            expect(portal.locator(selector)).to_have_text(name)
            portal.locator('#email-input').fill('dev@example.com')
            portal.locator('#password-input').fill('test123')
            portal.locator(selector).click()
            expect(portal.locator('#login-result')).to_have_text('Login successful')
        portal.goto('http://127.0.0.1:5173/developer?recovery_demo=1')
        page.goto('http://127.0.0.1:5174/')
        page.get_by_role('button', name='Saved Workflows', exact=True).click()
        card = page.locator('.workflow-card').filter(has=page.get_by_role('heading', name=workflow['plan']['workflow_name'], exact=True)).first
        expect(card).to_be_visible()
        # Clear all tab state to prove the library is fetched from SQLite.
        page.evaluate('sessionStorage.clear(); localStorage.clear()')
        page.reload()
        page.get_by_role('button', name='Saved Workflows', exact=True).click()
        expect(card).to_be_visible()
        card.get_by_role('button', name='Load Workflow', exact=True).click()
        expect(page.get_by_role('region', name='Loaded workflow')).to_contain_text('#login-button')
        if active:
            response = client.post(f"/runs/{active[0]['id']}/resume")
            run_id = response.json()['id']
            page.get_by_role('button', name='Runs', exact=True).click()
            page.get_by_role('button', name='View run').first.click()
        else:
            with page.expect_response(lambda response: response.url.endswith('/runs') and response.request.method == 'POST') as response:
                card.get_by_role('button', name='Run Recovery Demo', exact=True).click()
            assert response.value.status == 200
            run_id = response.value.json()['id']
        panel = page.get_by_role('region', name='Adaptive Recovery Agent')
        # A clean isolated validation DB gives an initial approval, then reusable memory.
        expect(panel.get_by_role('heading', name='Website change detected')).to_be_visible(timeout=45000)
        expect(panel).to_contain_text('Sign In')
        before = client.get(f'/runs/{run_id}').json()
        assert before['processed'] == 0
        page.screenshot(path=str(ARTIFACTS / 'recovery-approval.png'), full_page=True)
        panel.get_by_role('button', name='Approve & Continue').click()
        first = wait_run(run_id)
        assert first['ERROR'] == 0
        assert (first['PASS'], first['FAIL']) == (5, 1), 'Changed mode must preserve all login outcomes'
        entries = client.get(f'/runs/{run_id}/recoveries').json()
        assert all(e['state'] == 'recovered' for e in entries)
        assert any(e['method'] == 'Approved memory' for e in entries)
        calls = client.get('/metrics').json()['total_ai_calls']
        repeat = client.post('/runs', json={'workflow_id': workflow['id'], 'dataset_id': workflow['dataset_id'], 'recovery_demo': True}).json()
        second = wait_run(repeat['id'])
        assert second['ERROR'] == 0 and second['PASS'] == first['PASS']
        assert client.get('/metrics').json()['total_ai_calls'] == calls
        assert all(e['method'] == 'Approved memory' for e in client.get(f"/runs/{repeat['id']}/recoveries").json())
        page.get_by_role('button', name='Saved Workflows', exact=True).click()
        with page.expect_response(lambda response: response.url.endswith('/runs') and response.request.method == 'POST') as normal_response:
            card.get_by_role('button', name='Run Saved Workflow', exact=True).click()
        normal = wait_run(normal_response.value.json()['id'])
        assert (normal['PASS'], normal['FAIL'], normal['ERROR']) == (5, 1, 0)
        assert client.get(f"/runs/{normal['id']}/recoveries").json() == []
        book = load_workbook(io.BytesIO(client.get(f'/runs/{run_id}/download').content))
        assert 'Recoveries' in book.sheetnames
        assert client.post(f'/runs/{run_id}/debug/package').status_code == 200
        with zipfile.ZipFile(io.BytesIO(client.get(f'/runs/{run_id}/debug/package/download').content)) as archive:
            logs = json.loads(archive.read('runtime_logs.json'))
            assert logs['recoveries']
            assert 'test123' not in json.dumps(logs) and 'dev@example.com' not in json.dumps(logs)
        page.screenshot(path=str(ARTIFACTS / 'recovery-complete.png'), full_page=True)
        print('PASS: both query modes, identical login, saved workflow after refresh and load, real browser Login -> Sign In, approval, resume, all rows, memory reuse, no repeated AI calls, Excel and Bob audit.')
    finally:
        browser.close()

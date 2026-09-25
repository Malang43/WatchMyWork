"""Exercise the actual extension + app + API + executor. Start all three servers first.

Default: real Nemotron inference. --manual: explicitly test the labeled manual-review path.
All data used here is synthetic. This script never reads or prints .env.
"""
import argparse
import io
import json
import re
import sys
import tempfile
import time
from pathlib import Path
import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'test-results'
ARTIFACTS.mkdir(exist_ok=True)
args = argparse.ArgumentParser()
args.add_argument('--manual', action='store_true')
args = args.parse_args()
client = httpx.Client(base_url='http://127.0.0.1:8000', headers={'X-WatchMyWork': 'local-demo'}, timeout=100)

def check(response):
    if not response.is_success:
        raise AssertionError(f'Local API failed with HTTP {response.status_code}')
    return response.json()

def wait_run(run_id):
    until = time.monotonic() + 120
    while time.monotonic() < until:
        run = check(client.get(f'/runs/{run_id}'))
        if run['state'] == 'completed':
            return run
        if run['state'] in ('paused', 'stopped'):
            raise AssertionError(f"Run unexpectedly {run['state']}: {run['error']}")
        time.sleep(.3)
    raise AssertionError('Execution did not complete in 120 seconds')

for port in [5173, 5174, 8000]:
    response = httpx.get(f'http://127.0.0.1:{port}/' + ('health' if port == 8000 else ''))
    assert response.is_success, f'Start the server on port {port}'
baseline = check(client.get('/metrics'))['total_ai_calls']
with sync_playwright() as pw:
    extension = str(ROOT / 'extension')
    profile = tempfile.mkdtemp(prefix='watchmywork-e2e-')
    context = pw.chromium.launch_persistent_context(profile, channel='chromium', headless=True, args=[f'--disable-extensions-except={extension}', f'--load-extension={extension}'], viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
    app = context.new_page()
    errors = []
    app.on('pageerror', lambda error: errors.append(str(error)))
    try:
        app.goto('http://127.0.0.1:5174/')
        expect(app.get_by_role('heading', name='Show one browser test. Generate and run the rest.')).to_be_visible()
        app.screenshot(path=str(ARTIFACTS / 'dashboard.png'), full_page=True)
        app.get_by_role('button', name='New Workflow', exact=True).last.click()
        expect(app.get_by_role('heading', name='Show us how you work.', exact=True)).to_be_visible()
        expect(app.get_by_label('Upload spreadsheet')).to_be_visible()
        with app.expect_response(lambda r: r.url.endswith('/datasets') and r.request.method == 'POST') as upload:
            app.get_by_label('Upload spreadsheet').set_input_files(str(ROOT / 'sample-data' / 'tracking-demo.xlsx'))
        dataset = upload.value.json()
        assert dataset['record_count'] == 20
        app.get_by_role('button', name='Continue to Teach').click()
        app.get_by_role('button', name='Start Recording').click()
        mock = context.new_page()
        mock.on('pageerror', lambda error: errors.append(str(error)))
        mock.goto('http://127.0.0.1:5173/')
        expect(mock.locator('#watchmywork-recorder')).to_have_attribute('data-recording', re.compile(r'^[a-f0-9]{32}$'), timeout=15000)
        mock.get_by_label('Tracking ID', exact=True).fill('PK100001')
        mock.get_by_role('button', name='Check Status').click()
        expect(mock.locator('#tracking-result .status')).to_have_text('Delivered')
        expect(mock.locator('#watchmywork-recorder')).to_contain_text(re.compile(r'^\S Result saved'), timeout=15000)
        app.bring_to_front()
        with app.expect_response(lambda r: '/demos/' in r.url and r.url.endswith('/stop')) as stopped:
            app.get_by_role('button', name='Stop Recording').click()
        assert stopped.value.json()['state'] == 'complete'
        recorded = stopped.value.json()
        assert {'navigate', 'fill', 'click', 'extract'}.issubset({e['action'] for e in recorded['events']})
        print('PASS: spreadsheet upload and real extension semantic recording', flush=True)
        if args.manual:
            # Exercise the UI manual-review button through an injected API failure,
            # while the workflow itself still comes from the production manual endpoint.
            app.route('**/workflows/infer', lambda route: route.fulfill(status=503, content_type='application/json', body=json.dumps({'detail': 'Simulated free-model outage for manual-path testing.'})))
            app.get_by_role('button', name='Analyze Workflow').click()
            expect(app.get_by_role('alert')).to_contain_text('Simulated')
            with app.expect_response(lambda r: r.url.endswith('/workflows/manual')) as inferred:
                app.get_by_role('button', name='Review demonstrated steps manually').click()
        else:
            for attempt in range(3):
                with app.expect_response(lambda r: r.url.endswith('/workflows/infer'), timeout=135000) as inferred:
                    app.get_by_role('button', name='Analyze Workflow' if attempt == 0 else 'Retry').click()
                if inferred.value.status == 200:
                    break
                print(f'Inference attempt {attempt + 1}: friendly error received; ' + ('retrying the saved demonstration' if attempt < 2 else 'stopping test'), flush=True)
        response = inferred.value
        if response.status != 200:
            message = response.json().get('detail', 'Inference unavailable')
            print('LIVE INFERENCE UNAVAILABLE:', message, flush=True)
            raise AssertionError('The live inference request did not produce a workflow')
        workflow = response.json()
        calls_after_inference = check(client.get('/metrics'))['total_ai_calls']
        wid = workflow['id']
        assert workflow['source'] == ('manual' if args.manual else 'nemotron')
        expect(app.get_by_role('heading', name='Review Learned Test')).to_be_visible()
        assert client.post('/runs', json={'workflow_id': wid, 'dataset_id': dataset['id']}).status_code == 409
        print('PASS: validated inference and unconfirmed execution blocked', flush=True)
        app.get_by_role('button', name='Confirm Workflow', exact=True).click()
        with app.expect_response(lambda r: r.url.endswith('/runs') and r.request.method == 'POST') as started:
            app.get_by_role('button', name='Run Workflow').click()
        run_id = started.value.json()['id']
        assert check(client.post(f'/runs/{run_id}/pause'))['state'] == 'paused'
        expect(app.get_by_role('heading', name='Workflow paused')).to_be_visible(timeout=15000)
        with app.expect_response(lambda r: r.url.endswith(f'/runs/{run_id}/resume')) as resumed:
            app.get_by_role('button', name='Resume', exact=True).click()
        assert resumed.value.status == 200
        result = wait_run(run_id)
        assert result['processed'] == 20 and result['successful'] == 4 and result['manual_review'] == 16 and result['failed'] == 0, result
        expect(app.get_by_role('heading', name='Workflow Complete')).to_be_visible(timeout=15000)
        with app.expect_download() as download:
            app.get_by_role('link', name='Download Updated Excel').click()
        path = ARTIFACTS / 'updated-tracking-demo.xlsx'
        download.value.save_as(path)
        book = load_workbook(path)
        assert [book.active.cell(i, 2).value for i in range(2, 6)] == ['Delivered', 'Pending', 'In Transit', 'Returned']
        assert book['Exceptions'].max_row == 17
        print('PASS: approval, pause/resume, 20-row browser execution and Excel export', flush=True)
        app.get_by_role('button', name='Review Exceptions').click()
        expect(app.get_by_role('heading', name='Exception Queue')).to_be_visible()
        app.get_by_label('Reviewed status for row 6', exact=True).fill('Not Found — reviewed')
        app.get_by_role('button', name='Resolve', exact=True).first.click()
        expect(app.get_by_label('Reviewed status for row 6', exact=True)).to_have_count(0)
        app.get_by_role('button', name='Save Workflow', exact=True).click()
        expect(app.get_by_role('button', name='Workflow Saved')).to_be_visible()
        app.screenshot(path=str(ARTIFACTS / 'results.png'), full_page=True)
        app.get_by_role('button', name='Saved Workflows', exact=True).click()
        app.get_by_role('button', name='Use on another spreadsheet').first.click()
        expect(app.get_by_role('heading', name='Same workflow. Fresh data.')).to_be_visible()
        expect(app.get_by_label('Upload spreadsheet')).to_be_visible()
        with app.expect_response(lambda r: r.url.endswith('/datasets') and r.request.method == 'POST') as reuploaded:
            app.get_by_label('Upload spreadsheet').set_input_files(str(ROOT / 'sample-data' / 'tracking-demo.csv'))
        new_dataset = reuploaded.value.json()
        app.get_by_role('button', name='Review Workflow').click()
        with app.expect_response(lambda r: r.url.endswith('/runs') and r.request.method == 'POST') as rerun:
            app.get_by_role('button', name='Run Workflow').click()
        fresh_run = rerun.value.json()
        result = wait_run(fresh_run['id'])
        assert result['processed'] == 20 and result['successful'] == 4
        metrics = check(client.get('/metrics'))
        assert metrics['total_ai_calls'] == calls_after_inference, 'Execution/reuse must not call the LLM'
        # Verify stopping a third run leaves an exportable checkpoint.
        third = check(client.post('/runs', json={'workflow_id': wid, 'dataset_id': new_dataset['id']}))
        stopped_run = check(client.post(f"/runs/{third['id']}/stop"))
        assert stopped_run['state'] == 'stopped'
        assert client.get(f"/runs/{third['id']}/download").is_success
        app.get_by_role('button', name='Dashboard', exact=True).click()
        app.set_viewport_size({'width': 390, 'height': 844})
        app.screenshot(path=str(ARTIFACTS / 'mobile.png'), full_page=True)
        assert app.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile page overflows'
        assert errors == [], errors
        print('PASS: exception resolution, saved-workflow reuse, stop, zero per-row AI calls, mobile layout', flush=True)
        (ARTIFACTS / ('e2e-manual-summary.json' if args.manual else 'e2e-summary.json')).write_text(json.dumps({'mode': 'manual' if args.manual else 'live_nemotron', 'workflow_id': wid, 'run_id': run_id, 'records': 20, 'initial_successful': 4, 'initial_manual_review': 16, 'failed': 0, 'ai_calls_added': calls_after_inference - baseline, 'browser_errors': errors}, indent=2), encoding='utf-8')
    except Exception:
        app.screenshot(path=str(ARTIFACTS / 'failure.png'), full_page=True)
        raise
    finally:
        if 'dataset' in locals():
            for demo in check(client.get('/demos', params={'dataset_id': dataset['id']})):
                if demo['state'] == 'recording':
                    client.post(f"/demos/{demo['id']}/stop")
        context.close()

import io
import pytest
from openpyxl import load_workbook
from backend import storage, executor
from backend.schema import Workflow, demonstrated_plan
from backend.sheets import export_sheet, read_sheet

def test_upload_metadata_and_private_storage(client, dataset):
    assert dataset['record_count'] == 4
    assert dataset['columns'] == ['Tracking ID', 'Status']
    assert 'rows' not in dataset
    assert client.get(f"/datasets/{dataset['id']}/rows/0").json()['Tracking ID'] == 'PK100001'
    assert client.get(f"/datasets/{dataset['id']}/rows/99").status_code == 404

@pytest.mark.parametrize('name,content', [('bad.csv', b'ID,ID\nx,y'), ('empty.csv', b'ID,Status'), ('bad.xlsx', b'broken'), ('bad.txt', b'ID,Status\nx,y'), ('wide.csv', b'ID,Status\nx,y,z'), ('headers.csv', b',Status\nx,y')])
def test_invalid_upload(client, name, content):
    assert client.post('/datasets', files={'file': (name, content)}).status_code == 422

def test_origin_header_guard(client):
    response = client.post('/datasets', headers={'X-WatchMyWork': ''}, files={'file': ('x.csv', b'ID,Status\nx,y')})
    assert response.status_code == 403
    assert client.get('/health', headers={'Host': 'attacker.invalid'}).status_code == 403

def test_cannot_execute_or_save_unconfirmed(client, workflow, dataset):
    assert client.post('/runs', json={'workflow_id': workflow['id'], 'dataset_id': dataset['id']}).status_code == 409
    assert client.post(f"/workflows/{workflow['id']}/save").status_code == 409

def test_approval_checkpoint_export_reuse(client, workflow, dataset, monkeypatch):
    monkeypatch.setattr(executor, 'start', lambda run_id: None)
    wid = workflow['id']
    assert client.post(f'/workflows/{wid}/confirm').status_code == 200
    assert client.post(f'/workflows/{wid}/save').json()['saved']
    run = client.post('/runs', json={'workflow_id': wid, 'dataset_id': dataset['id']}).json()
    rid = run['id']
    assert run['processed'] == 1 and run['successful'] == 1
    assert storage.results(rid)[0]['source'] == 'demonstration'
    assert client.post(f'/runs/{rid}/pause').json()['state'] == 'paused'
    assert client.post(f'/runs/{rid}/resume').json()['state'] == 'running'
    storage.checkpoint(rid, 1, 'successful', value='Pending')
    storage.checkpoint(rid, 2, 'manual_review', reason='Duplicate ID')
    assert client.post(f'/runs/{rid}/stop').json()['processed'] == 3
    assert client.post(f'/runs/{rid}/resume').status_code == 409
    exported = client.get(f'/runs/{rid}/download')
    book = load_workbook(io.BytesIO(exported.content))
    assert book.worksheets[0]['B2'].value == 'Delivered'
    assert book.worksheets[0]['B3'].value == 'Pending'
    assert book['Exceptions'].max_row == 2
    assert client.put(f'/runs/{rid}/exceptions/2', json={'value': 'Verified duplicate'}).status_code == 200
    assert client.get(f'/runs/{rid}/exceptions').json() == []
    fresh = client.post('/datasets', files={'file': ('fresh.csv', b'Tracking ID,Status\nPK100004,')}).json()
    rerun = client.post('/runs', json={'workflow_id': wid, 'dataset_id': fresh['id']}).json()
    assert rerun['processed'] == 0  # Never reuse demonstration results on a different spreadsheet.
    metrics = client.get('/metrics').json()
    assert 'PK100001' not in str(metrics) and 'Verified duplicate' not in str(metrics)

def test_edits_revoke_approval(client, workflow):
    wid = workflow['id']
    client.post(f'/workflows/{wid}/confirm')
    plan = workflow['plan'] | {'workflow_name': 'Reviewed tracking lookup'}
    changed = client.put(f'/workflows/{wid}', json=plan).json()
    assert changed['confirmed'] is False and changed['human_corrections'] == 1

@pytest.mark.parametrize('mutation', ['url', 'selector', 'action', 'value', 'column', 'extra'])
def test_rejects_unsafe_ai_plans(mutation):
    plan = demonstrated_plan('Tracking ID', 'Status').model_dump()
    if mutation == 'url': plan['url'] = 'https://example.com/'
    elif mutation == 'selector': plan['steps'][0]['target'] = 'body'
    elif mutation == 'action': plan['steps'][1]['action'] = 'evaluate'
    elif mutation == 'value': plan['steps'][0]['value'] = 'PK100001'
    elif mutation == 'column': plan['steps'][3]['column'] = 'Tracking ID'
    else: plan['code'] = 'alert(1)'
    with pytest.raises(ValueError): Workflow.model_validate(plan)

def test_requires_real_complete_demonstration(client, dataset):
    request = {'dataset_id': dataset['id'], 'input_column': 'Tracking ID', 'destination_column': 'Status'}
    assert client.post('/workflows/infer', json=request).status_code == 422
    demo = client.post('/demos', json={**request, 'row_index': 0}).json()
    assert client.post('/demos', json={**request, 'row_index': 1}).status_code == 409
    client.post(f"/demos/{demo['id']}/events", json={'events': [{'action': 'extract', 'target': '#tracking-result .status', 'text': 'Delivered'}]})
    assert client.post(f"/demos/{demo['id']}/stop").json()['state'] == 'incomplete'
    assert client.post('/workflows/manual', json=request).status_code == 422

def test_crash_recovery_preserves_checkpoint(client, workflow, dataset, monkeypatch):
    monkeypatch.setattr(executor, 'start', lambda _: None)
    client.post(f"/workflows/{workflow['id']}/confirm")
    run = client.post('/runs', json={'workflow_id': workflow['id'], 'dataset_id': dataset['id']}).json()
    storage.init()
    assert client.get(f"/runs/{run['id']}").json()['state'] == 'paused'
    assert client.get(f"/runs/{run['id']}").json()['successful'] == 1

def test_export_formula_strings_are_not_executable():
    columns, rows = read_sheet(b'ID,Status\n=HYPERLINK("bad"),', 'test.csv')
    content = export_sheet({'columns': columns, 'rows': rows}, 'Status', [])
    book = load_workbook(io.BytesIO(content))
    assert book.active['A2'].data_type == 's'

@pytest.mark.parametrize('url,allowed', [('http://127.0.0.1:5173/', True), ('http://127.0.0.1:5173/src/main.tsx', True), ('https://example.com/', False), ('http://127.0.0.1:8000/', False), ('http://evil@127.0.0.1:5173/', False)])
def test_executor_network_allowlist(url, allowed):
    assert executor.allowed_url(url) is allowed

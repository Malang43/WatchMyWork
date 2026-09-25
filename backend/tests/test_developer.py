import io
import json
from pathlib import Path
from openpyxl import load_workbook
from backend.schema import Mapping, Workflow, demonstrated_plan
from backend.inference import compile_demonstration, evidence
from backend.sheets import export_sheet
from backend.developer import generate
import pytest

MAPPING = dict(mode='developer', input_columns=['Email', 'Password'], destination_column='Actual Result', expected_column='Expected Result', status_column='Status')

def test_developer_schema_and_codegen():
    plan = demonstrated_plan(Mapping(**MAPPING), 'Actual Result')
    assert plan.steps[0].target == '#email-input'
    assert plan.steps[1].value == '{{row.Password}}'
    assert plan.url.endswith('/developer')
    code = generate(plan.model_dump())
    compile(code, 'generated_test.spec.py', 'exec')
    assert 'for index, row in enumerate(cases, 1)' in code
    for field, value in [('url', 'http://127.0.0.1:5173/'), ('status_column', 'Email')]:
        with pytest.raises(ValueError):
            Workflow.model_validate({**plan.model_dump(), field: value})

def test_developer_record_compile_and_privacy(client):
    data = client.post('/datasets', files={'file': ('login.csv', Path('sample-data/developer-tests.csv').read_bytes(), 'text/csv')}).json()
    request = dict(dataset_id=data['id'], **MAPPING)
    demo = client.post('/demos', json={**request, 'row_index': 0}).json()
    events = [dict(action='fill', target='#email-input', value='dev@example.com'), dict(action='fill', target='#password-input', value='test123'), dict(action='click', target='#login-button'), dict(action='wait', target='#login-result'), dict(action='extract', target='#login-result', text='Login successful')]
    for event in events: event['url'] = 'http://127.0.0.1:5173/developer'
    response = client.post(f"/demos/{demo['id']}/events", json={'events': events})
    assert response.json()['state'] == 'complete'
    demos = client.get('/demos', params={'dataset_id': data['id']}).json()
    safe = json.dumps(evidence(demos, Mapping(**MAPPING), 'Actual Result'))
    assert 'test123' not in safe and 'dev@example.com' not in safe and 'Login successful' not in safe
    plan = compile_demonstration(demos, Mapping(**MAPPING), 'Actual Result')
    assert plan.mode == 'developer'
    workflow = client.post('/workflows/manual', json=request).json()
    assert client.get(f"/workflows/{workflow['id']}/generated-test").status_code == 409
    client.post(f"/workflows/{workflow['id']}/confirm")
    assert client.get(f"/workflows/{workflow['id']}/generated-test").status_code == 200
    assert client.post(f"/demos/{demo['id']}/events", json={'events': events}).json()['recorded'] == len(events)

def test_failed_actual_is_preserved():
    dataset = {'columns': ['Expected', 'Actual', 'Status'], 'rows': [{'Expected': 'yes', 'Actual': '', 'Status': ''}]}
    result = [{'row_index': 0, 'status': 'FAIL', 'value': ' No ', 'reason': 'mismatch'}]
    book = load_workbook(io.BytesIO(export_sheet(dataset, 'Actual', result, 'Status')))
    assert book.active.cell(2, 2).value == ' No '
    assert book.active.cell(2, 3).value == 'FAIL'

def test_runtime_error_continues_and_blank_duplicate_execute(client, monkeypatch):
    from backend import storage, executor
    from backend.tests.test_executor import browser_fixture
    from playwright.sync_api import TimeoutError
    page = browser_fixture(monkeypatch)
    page.url = 'http://127.0.0.1:5173/developer'
    page.goto.side_effect = [TimeoutError('private diagnostic'), TimeoutError('private diagnostic'), None, None]
    page.locator.return_value.inner_text.return_value = ' Login successful '
    row = {'Email': '', 'Password': '', 'Expected Result': 'login SUCCESSFUL', 'Actual Result': '', 'Status': ''}
    data = storage.put('dataset', {'id': storage.uid(), 'columns': list(row), 'rows': [dict(row), dict(row), dict(row)]})
    run = storage.put('run', {'id': storage.uid(), 'dataset_id': data['id'], 'plan': demonstrated_plan(Mapping(**MAPPING), 'Actual Result').model_dump(), 'state': 'running', 'retries': 0})
    executor.execute(run['id'])
    outcomes = storage.results(run['id'])
    assert [r['status'] for r in outcomes] == ['ERROR', 'PASS', 'PASS']
    assert outcomes[1]['value'] == ' Login successful '
    assert storage.get('run', run['id'])['state'] == 'completed'
    assert 'private diagnostic' not in json.dumps(outcomes)

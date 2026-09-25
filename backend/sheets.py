import csv
import io
import zipfile
from openpyxl import Workbook, load_workbook

MAX_ROWS = 10000
MAX_COLUMNS = 100

def read_sheet(content: bytes, filename: str):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('File must be under 10 MB')
    if filename.lower().endswith('.csv'):
        rows = list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
    elif filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Workbook expands beyond the 50 MB limit')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError('Maximum 10,000 rows and 100 columns')
            rows = [[str(c.value) if c.value is not None else '' for c in row] for row in sheet.iter_rows()]
        finally:
            book.close()
    else:
        raise ValueError('Choose a .csv or .xlsx file')
    if len(rows) < 2 or len(rows) > MAX_ROWS + 1:
        raise ValueError('Include a header and 1–10,000 data rows')
    columns = [str(c).strip() for c in rows[0]]
    if not columns or len(columns) > MAX_COLUMNS or any(not c or len(c) > 100 for c in columns) or len(set(columns)) != len(columns):
        raise ValueError('Column names must be unique, nonblank, and at most 100 characters')
    if any(len(row) > len(columns) for row in rows[1:]):
        raise ValueError('A row has more values than the header')
    values = [{col: str(row[i]) if i < len(row) else '' for i, col in enumerate(columns)} for row in rows[1:]]
    if any(len(v) > 10000 for row in values for v in row.values()):
        raise ValueError('A cell exceeds the 10,000 character limit')
    return columns, values

def export_sheet(dataset, output_column, results, status_column=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Updated Data'
    columns = dataset['columns']
    sheet.append(columns)
    by_index = {r['row_index']: r for r in results}
    for index, row in enumerate(dataset['rows']):
        values = dict(row)
        result = by_index.get(index)
        if result and result['status'] in ('successful', 'PASS', 'FAIL'):
            values[output_column] = result['value']
        if result and status_column:
            values[status_column] = result['status']
        sheet.append([values.get(c, '') for c in columns])
    # Always export literal strings, never executable spreadsheet formulas.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 24
    audit = book.create_sheet('Exceptions')
    audit.append(['Spreadsheet row', 'Outcome', 'Reason'])
    for result in results:
        if result['status'] != 'successful':
            audit.append([result['row_index'] + 2, result['status'], result['reason']])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
